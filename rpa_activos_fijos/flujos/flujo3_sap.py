"""
flujos/flujo3_sap.py — FLUJO 3: SAP, activo por activo.

CAMBIO DE NEGOCIO (2026-10-04): NO hay carga masiva. Cada fila del Excel
(ya validado en el Flujo 2) es UN activo, y se trabaja con la transacción
individual que corresponda (BRP - Creación: AS01), llenando el formulario de
SAP con los valores de esa fila. SAP NO recibe el archivo.

El formulario es CONFIGURACIÓN (config.py, SAP_FORMULARIOS_POR_CASO): una
lista ordenada de pasos (campo / casilla / Enter / pestaña) con la columna
del Excel y el XPath de cada campo.

CASO ESPECIAL "dos archivos" (BRP, 2026-10-10): la solicitud trae el Excel
de modificación y el de creación: PASO 1 todas las filas de modificación en
AS02 (actualizar el número de inventario) y PASO 2 todas las de creación en
AS01, igual que una creación normal.

Por cada fila:
  0. (desde la 2ª fila) recargar la página de SAP: SAP vuelve al inicio con
     la página nueva y estable (sin esto, la 2ª fila fallaba con elementos
     "vencidos", 1ª prueba real 2026-10-05).
  1. /n + transacción.
  2. Llenar el formulario (celdas vacías no se tocan).
  3a. Si guardar_permitido(): Guardar -> leer el mensaje de la barra -> código
      del activo (ej. "El act.fj. 7129560 0 se ha creado" -> 7129560) o el
      mensaje de error. Va a solicitud.resultados_sap[fila].
  3b. Si NO (interruptor apagado o MODO PRUEBAS): se DETIENE, pide
      "Continuar" en la ventana del bot y sale SIN guardar.
  Si algo falla en una fila: se anota el mensaje, se registra en el log y se
  CONTINÚA con la siguiente fila.

Seguridad (no hay SAP de pruebas): ver guardar_permitido(). Además, nunca se
guarda si no está configurado dónde leer el mensaje de resultado.
"""

import os
import re
from datetime import date, datetime

from openpyxl import load_workbook

from config import (
    CASO_ESPECIAL_CREAR_SI_FALLA_MODIFICACION,
    CASOS_DOS_ARCHIVOS,
    MODO_PRUEBAS_REEMPLAZO,
    MODO_PRUEBAS_SOLO_SAP,
    OUTPUT_DIR,
    SAP_ESPERA_ENTRE_CLICS_SEG,
    SAP_FORMATO_FECHA,
    SAP_FORMULARIOS_POR_CASO,
    SAP_GUARDAR_REAL,
    SAP_REGEX_EXITO_POR_TRANSACCION,
    SAP_SEPARADOR_DECIMAL,
    SAP_TRANSACCION_POR_CASO,
    SAP_XPATH_BOTON_ATRAS,
    SAP_XPATH_BOTON_GUARDAR,
    SAP_XPATH_CONFIRMAR_SALIR_SIN_GUARDAR,
    SAP_XPATH_MENSAJE_ESTADO,
)
from core.exceptions import SapError
from validacion import router
from validacion.base_validador import es_vacia

ENCABEZADO_CODIGO_SAP = "Código SAP"

# Para el resumen del caso: "SAP: 3 creado(s), 1 con error".
_VERBO_EXITO = {"AS01": "creado(s)", "AS02": "modificado(s)", "AS06": "borrado(s)"}


def guardar_permitido():
    """
    True solo si se puede presionar "Guardar" en SAP (crea / modifica / borra
    activos REALES). Hace falta SAP_GUARDAR_REAL = True Y NO estar en ningún
    modo pruebas (REEMPLAZO ni SOLO_SAP): en modo pruebas guardar está
    PROHIBIDO siempre.
    """
    return (
        bool(SAP_GUARDAR_REAL)
        and not MODO_PRUEBAS_REEMPLAZO
        and not MODO_PRUEBAS_SOLO_SAP
    )


def valor_para_sap(valor):
    """Convierte el valor de una celda al texto que se escribe en SAP."""
    if isinstance(valor, (datetime, date)):
        return valor.strftime(SAP_FORMATO_FECHA)
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor))
        return str(valor).replace(".", SAP_SEPARADOR_DECIMAL)
    return str(valor).strip()


def cargar_a_sap(sap, solicitud, logger=None, esperar_continuar=None):
    """
    Trabaja en SAP cada fila de la(s) plantilla(s) validada(s).

    Normal: un Excel -> su transacción (BRP - Creación: AS01).
    CASO ESPECIAL "dos archivos" (solicitud.excel_modificacion): PASO 1 todas
    las filas del Excel de modificación (AS02) y PASO 2 todas las del de
    creación (AS01). Si el PASO 1 tiene filas con error y
    CASO_ESPECIAL_CREAR_SI_FALLA_MODIFICACION = False, el PASO 2 NO se hace.

    Args:
        sap (SapWebGui): sesión de SAP (se abre aquí la primera vez).
        solicitud (Solicitud): caso con su(s) plantilla(s) ya validada(s). Se
            llenan solicitud.resultados_sap (y resultados_sap_modificacion).
        logger: logger opcional.
        esperar_continuar (callable, opcional): función que BLOQUEA hasta que
            la persona presione "Continuar" en la ventana del bot. Recibe un
            texto para mostrar. Si es None, no se espera (pruebas).

    Returns:
        str: texto del paso en el que quedó (para el resumen del caso).

    Raises:
        SapError: si falta configuración (transacción, mensaje de éxito para
            guardar), o si en el caso especial falló la modificación y por
            eso no se hizo la creación.
    """
    case_id = solicitud.case_id

    # Archivos a trabajar, EN ORDEN: (tipo, accion, ruta, resultados, etiqueta)
    trabajos = []
    if solicitud.excel_modificacion:
        tipo_previo, accion_previa = CASOS_DOS_ARCHIVOS[(solicitud.tipo, solicitud.accion)]
        trabajos.append((tipo_previo, accion_previa, solicitud.excel_modificacion,
                         solicitud.resultados_sap_modificacion, "PASO 1 (modificación)"))
    trabajos.append((solicitud.tipo, solicitud.accion, solicitud.excel_path,
                     solicitud.resultados_sap,
                     "PASO 2 (creación)" if solicitud.excel_modificacion else None))

    # 1) Revisar la configuración de TODOS los archivos antes de tocar SAP.
    guardar = guardar_permitido()
    if guardar and not SAP_XPATH_MENSAJE_ESTADO:
        raise SapError(
            "Se pidió guardar en SAP, pero SAP_XPATH_MENSAJE_ESTADO está vacío: "
            "sin leer el mensaje no se puede confirmar si el activo se creó. "
            "No se guardó nada."
        )
    planes = []
    for tipo, accion, ruta, resultados, etiqueta in trabajos:
        transaccion = SAP_TRANSACCION_POR_CASO.get((tipo, accion))
        if not transaccion:
            raise SapError(
                f"Caso {case_id}: (tipo={tipo!r}, accion={accion!r}) "
                "no tiene transacción de SAP configurada. Revisa "
                "SAP_TRANSACCION_POR_CASO en config.py."
            )
        formulario = SAP_FORMULARIOS_POR_CASO.get((tipo, accion))
        if not formulario:
            if logger:
                logger.warning(
                    "Caso %s: PENDIENTE: el formulario de %s aún no está configurado; "
                    "no se entra a SAP. Plantilla validada: %s",
                    case_id, transaccion, ruta,
                )
            return f"Pendiente de SAP ({transaccion} sin configurar)"
        if guardar and transaccion not in SAP_REGEX_EXITO_POR_TRANSACCION:
            raise SapError(
                f"Se pidió guardar en {transaccion}, pero su mensaje de éxito no está "
                "configurado (SAP_REGEX_EXITO_POR_TRANSACCION): no se podría saber si "
                "salió bien. No se guardó nada."
            )
        planes.append((tipo, accion, ruta, resultados, etiqueta, transaccion, formulario))

    # 2) Trabajar cada archivo, en orden.
    sap.asegurar_sesion()
    textos = []
    ya_hubo_filas = False
    for tipo, accion, ruta, resultados, etiqueta, transaccion, formulario in planes:
        if etiqueta and logger:
            logger.info("Caso %s: === %s en %s ===", case_id, etiqueta, transaccion)
        filas, errores = _trabajar_archivo(
            sap, case_id, tipo, accion, ruta, resultados, transaccion, formulario,
            guardar, ya_hubo_filas, logger, esperar_continuar,
        )
        ya_hubo_filas = ya_hubo_filas or filas > 0
        texto = _texto_resultado(transaccion, filas, errores, guardar)
        textos.append(f"{etiqueta} {transaccion}: {texto}" if etiqueta else texto)

        # CASO ESPECIAL: si falló la modificación, ¿se crea igual?
        if etiqueta and resultados is solicitud.resultados_sap_modificacion and errores:
            if not CASO_ESPECIAL_CREAR_SI_FALLA_MODIFICACION:
                raise SapError(
                    f"{etiqueta} {transaccion}: {errores} de {filas} fila(s) con error "
                    f"(filas {', '.join(str(n) for n, t in resultados.items() if t.startswith('ERROR'))}). "
                    "Por eso NO se hizo el PASO 2 (creación) "
                    "[CASO_ESPECIAL_CREAR_SI_FALLA_MODIFICACION = False]."
                )
            if logger:
                logger.warning(
                    "Caso %s: el PASO 1 tuvo %s fila(s) con error; se sigue con la "
                    "creación (CASO_ESPECIAL_CREAR_SI_FALLA_MODIFICACION = True).",
                    case_id, errores,
                )

    if guardar:
        escribir_columna_codigo_sap(solicitud, logger=logger)
    return " | ".join(textos)


def _trabajar_archivo(sap, case_id, tipo, accion, ruta, resultados, transaccion,
                      formulario, guardar, recargar_al_empezar, logger, esperar_continuar):
    """
    Trabaja en SAP las filas de UN Excel. Devuelve (filas, errores). Antes de
    cada fila (salvo la primera de toda la solicitud) recarga la página.
    """
    filas = router.resolver(tipo, accion, logger=logger).leer_filas(ruta)
    if logger:
        logger.info(
            "Caso %s: %s activo(s) para %s. Guardar en SAP: %s.",
            case_id, len(filas), transaccion,
            "SÍ" if guardar else "NO (se detiene antes de Guardar en cada activo)",
        )

    for indice, (numero, fila) in enumerate(filas):
        try:
            if indice > 0 or recargar_al_empezar:
                # Entre filas: recargar la página (F5) para empezar con SAP
                # limpio. Tras salir de la transacción, SAP queda redibujando
                # y los campos se "vencen" (1ª prueba real, 2026-10-05).
                sap.recargar()
            sap.ir_a_transaccion(transaccion)
            _llenar_formulario(sap, formulario, fila, case_id, numero, logger)
            if guardar:
                resultados[numero] = _guardar_y_leer(sap, transaccion, case_id, numero, logger)
            else:
                _supervisar(sap, case_id, numero, transaccion, logger, esperar_continuar)
        except SapError as e:
            detalle = _mensaje_sap_si_hay(sap)
            texto = f"ERROR: {detalle}" if detalle else f"ERROR: {e}"
            resultados[numero] = texto
            if logger:
                logger.error("Caso %s | fila %s: %s%s (se continúa con la siguiente)",
                             case_id, numero, e,
                             f" | mensaje de SAP: {detalle}" if detalle else "")
            _descartar_tras_error(sap, logger)

    errores = sum(1 for t in resultados.values() if t.startswith("ERROR"))
    return len(filas), errores


def _texto_resultado(transaccion, filas, errores, guardar):
    """Texto corto para el resumen del caso (igual que antes del caso especial)."""
    if guardar:
        exitosos = filas - errores
        return f"SAP: {exitosos} {_VERBO_EXITO.get(transaccion, 'OK')}, {errores} con error"
    if errores:
        return f"Revisado en SAP SIN guardar ({filas} fila(s), {errores} con error)"
    return f"Revisado en SAP SIN guardar ({filas} fila(s))"


def _llenar_formulario(sap, formulario, fila, case_id, numero, logger):
    """Ejecuta los pasos del formulario con los valores de UNA fila."""
    for paso in formulario:
        tipo = paso["tipo"]
        if tipo == "enter":
            sap.enter()
        elif tipo == "pestana":
            sap.doble_clic_lento(paso["xpath"], paso["nombre"], SAP_ESPERA_ENTRE_CLICS_SEG)
        elif tipo == "casilla":
            if not es_vacia(fila.get(paso["columna"])):
                sap.clic(paso["xpath"], f"la casilla '{paso['nombre']}'")
        elif tipo == "campo":
            valor = fila.get(paso["columna"])
            if es_vacia(valor):
                continue  # celda vacía: el campo no se toca
            sap.escribir(paso["xpath"], valor_para_sap(valor),
                         f"fila {numero}, {paso['columna']} ({paso['nombre']})")
        else:
            raise SapError(f"Paso de formulario desconocido en config.py: {paso!r}")


def _guardar_y_leer(sap, transaccion, case_id, numero, logger):
    """Guarda y devuelve el código del activo o el mensaje de error de SAP."""
    sap.clic(SAP_XPATH_BOTON_GUARDAR, "el botón Guardar")
    mensaje = sap.leer_texto(SAP_XPATH_MENSAJE_ESTADO, "el mensaje de SAP tras guardar")
    coincidencia = re.search(SAP_REGEX_EXITO_POR_TRANSACCION[transaccion], mensaje, re.IGNORECASE)
    if coincidencia:
        codigo = coincidencia.group(1)
        if logger:
            logger.info("Caso %s | fila %s: %s OK -> %s (%s)",
                        case_id, numero, transaccion, codigo, mensaje)
        return codigo
    if logger:
        logger.error("Caso %s | fila %s: %s no se completó: %s",
                     case_id, numero, transaccion, mensaje)
    return f"ERROR: {mensaje}"


def _supervisar(sap, case_id, numero, transaccion, logger, esperar_continuar):
    """Modo supervisado: se detiene antes de Guardar, espera "Continuar" y
    sale SIN guardar."""
    texto = (f"Caso {case_id}, fila {numero}: formulario de {transaccion} lleno. "
             "NO se guarda. Revisa SAP y presiona 'Continuar'.")
    if logger:
        logger.warning(texto)
    if esperar_continuar:
        esperar_continuar(texto)
    _salir_sin_guardar(sap, logger)


def _salir_sin_guardar(sap, logger):
    """
    Sale del formulario SIN guardar (receta de la usuaria): dos clics en
    "Atrás" y confirmar "salir sin guardar" en la ventana que aparece. SAP
    vuelve al inicio de la transacción.
    """
    try:
        sap.doble_clic_lento(SAP_XPATH_BOTON_ATRAS, "el botón Atrás", SAP_ESPERA_ENTRE_CLICS_SEG)
        confirmar = sap.existe(SAP_XPATH_CONFIRMAR_SALIR_SIN_GUARDAR, SAP_ESPERA_ENTRE_CLICS_SEG * 3)
        if confirmar is not None:
            confirmar.click()
    except Exception as e:
        raise SapError(f"No se pudo salir del formulario sin guardar: {e}")


def _descartar_tras_error(sap, logger):
    """
    Después de un error la pantalla puede estar en cualquier estado. NO se
    presiona "Atrás" (fuera del formulario podría llevar al menú o a cerrar
    sesión): solo se confirma "salir sin guardar" si esa ventana quedó
    abierta. La siguiente fila igual arranca con /n + transacción. Nunca lanza.
    """
    try:
        confirmar = sap.existe(SAP_XPATH_CONFIRMAR_SALIR_SIN_GUARDAR, 1)
        if confirmar is not None:
            confirmar.click()
    except Exception as e:
        if logger:
            logger.warning("Tras el error no se pudo cerrar la ventana de SAP (%s); "
                           "la siguiente fila reinicia con /n.", e)


def _mensaje_sap_si_hay(sap):
    """Texto de la barra de mensajes de SAP, si está configurada y visible."""
    if not SAP_XPATH_MENSAJE_ESTADO:
        return None
    try:
        elemento = sap.existe(SAP_XPATH_MENSAJE_ESTADO, 2)
        texto = (elemento.text or "").strip() if elemento is not None else ""
        return texto or None
    except Exception:
        return None


def escribir_columna_codigo_sap(solicitud, logger=None):
    """
    Agrega la columna "Código SAP" al Excel del USUARIO (en la primera
    columna libre a la derecha) con el resultado de cada fila, y lo guarda
    APARTE en salidas/ (el original no se modifica). Devuelve la ruta.
    """
    origen = solicitud.excel_path
    base, extension = os.path.splitext(os.path.basename(origen))
    destino = os.path.join(OUTPUT_DIR, f"{base}_RESPUESTA{extension}")

    libro = load_workbook(origen, keep_vba=extension.lower() == ".xlsm")
    hoja = libro.worksheets[0]
    columna = hoja.max_column + 1
    hoja.cell(row=1, column=columna, value=ENCABEZADO_CODIGO_SAP)
    for numero, texto in solicitud.resultados_sap.items():
        hoja.cell(row=numero, column=columna, value=texto)
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    libro.save(destino)

    if logger:
        logger.info("Caso %s: Excel con '%s' -> %s",
                    solicitud.case_id, ENCABEZADO_CODIGO_SAP, destino)
    return destino
