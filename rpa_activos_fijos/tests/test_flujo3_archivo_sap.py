"""
Pruebas de la preparación del archivo para SAP: copia temporal con el nombre
EXACTO que exige SAP ("CREAR (BRP).xlsm" para BRP - Creación) y SIN las
columnas que SAP no acepta (la AC de BRP - Creación).
La carpeta carga_sap/ se redirige a una temporal.
"""

import os
import shutil
import zipfile

import pytest
from openpyxl import load_workbook

from core.exceptions import SapError
from core.models import Solicitud
from flujos import flujo3_sap

CARPETA_BRP = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "plantillas",
    "BRP",
)
PLANTILLA_USUARIO = os.path.join(CARPETA_BRP, "Plantilla Creación Activos BRP usuario.xlsm")
PLANTILLA_SAP = os.path.join(CARPETA_BRP, "Plantilla Creación Activos BRP.xlsm")


@pytest.fixture
def carga_sap(tmp_path, monkeypatch):
    carpeta = tmp_path / "carga_sap"
    carpeta.mkdir()
    monkeypatch.setattr(flujo3_sap, "CARGA_SAP_DIR", str(carpeta))
    return carpeta


@pytest.fixture
def sin_quitar_columnas(monkeypatch):
    """Para las pruebas de la COPIA en sí (archivos de mentira, no Excel)."""
    monkeypatch.setattr(flujo3_sap, "COLUMNAS_QUITAR_ANTES_DE_SAP", {})
    monkeypatch.setattr(flujo3_sap, "ULTIMA_COLUMNA_PLANTILLA", {})


def solicitud_con_archivo(tmp_path, case_id, contenido, tipo="brp", accion="creacion"):
    ruta = tmp_path / f"{case_id}_{tipo}_{accion}.xlsm"
    ruta.write_bytes(contenido)
    return Solicitud(case_id=case_id, tipo=tipo, accion=accion, excel_path=str(ruta))


# -- La copia con el nombre exacto ------------------------------------------------

def test_copia_con_el_nombre_exacto_de_sap(tmp_path, carga_sap, sin_quitar_columnas):
    solicitud = solicitud_con_archivo(tmp_path, "PDA-7889", b"datos 7889")

    ruta = flujo3_sap.preparar_archivo_sap(solicitud)

    assert os.path.basename(ruta) == "CREAR (BRP).xlsm"
    assert open(ruta, "rb").read() == b"datos 7889"
    assert solicitud.archivo_sap == ruta
    # El archivo de la solicitud sigue ahí, intacto y con su nombre.
    assert open(solicitud.excel_path, "rb").read() == b"datos 7889"


def test_cada_solicitud_reemplaza_la_copia_anterior(tmp_path, carga_sap, sin_quitar_columnas):
    primera = solicitud_con_archivo(tmp_path, "PDA-7889", b"datos 7889")
    segunda = solicitud_con_archivo(tmp_path, "PDA-7890", b"datos 7890")

    flujo3_sap.preparar_archivo_sap(primera)
    ruta = flujo3_sap.preparar_archivo_sap(segunda)

    assert os.listdir(carga_sap) == ["CREAR (BRP).xlsm"]
    assert open(ruta, "rb").read() == b"datos 7890"


def test_si_falla_la_copia_no_queda_el_archivo_de_otra_solicitud(tmp_path, carga_sap, sin_quitar_columnas):
    primera = solicitud_con_archivo(tmp_path, "PDA-7889", b"datos 7889")
    flujo3_sap.preparar_archivo_sap(primera)

    rota = Solicitud(case_id="PDA-7890", tipo="brp", accion="creacion",
                     excel_path=str(tmp_path / "no_existe.xlsm"))
    with pytest.raises(SapError):
        flujo3_sap.preparar_archivo_sap(rota)

    # Lo importante: el archivo de la PDA-7889 ya no está listo para cargarse.
    assert os.listdir(carga_sap) == []


def test_combinacion_sin_nombre_configurado(tmp_path, carga_sap):
    solicitud = solicitud_con_archivo(tmp_path, "PDA-1", b"x", tipo="prj")
    with pytest.raises(SapError, match="NOMBRE_ARCHIVO_SAP"):
        flujo3_sap.preparar_archivo_sap(solicitud)
    assert os.listdir(carga_sap) == []


# -- Quitar la columna AC (BRP - Creación) -------------------------------------------

def excel_usuario(tmp_path, filas):
    """Plantilla del usuario (A..AC) con `filas` desde la fila 2."""
    ruta = tmp_path / "PDA-7889_brp_creacion.xlsm"
    shutil.copy(PLANTILLA_USUARIO, ruta)
    libro = load_workbook(ruta, keep_vba=True)
    for numero, fila in enumerate(filas, start=2):
        for letra, valor in fila.items():
            libro.worksheets[0][f"{letra}{numero}"] = valor
    libro.save(ruta)
    return Solicitud(case_id="PDA-7889", tipo="brp", accion="creacion", excel_path=str(ruta))


def encabezados(ruta):
    hoja = load_workbook(ruta).worksheets[0]
    return [c.value for c in hoja[1]]


FILA = dict(A="BRP01", B="1000", C=1, D="PORTATIL", F="INV-1", H="CC1", O="NUEVO",
            W="AGR", AB="P1", AC="PROVEEDOR SAS NIT 900123456")


def test_copia_para_sap_queda_igual_a_la_plantilla_de_sap(tmp_path, carga_sap):
    solicitud = excel_usuario(tmp_path, [FILA, dict(FILA, F="INV-2", AC="OTRO")])

    ruta = flujo3_sap.preparar_archivo_sap(solicitud)

    # Mismos encabezados, mismo orden que la plantilla que recibe SAP (A..AB).
    assert encabezados(ruta) == encabezados(PLANTILLA_SAP)
    hoja = load_workbook(ruta).worksheets[0]
    assert hoja.max_column == 28
    # Los datos de A..AB se conservan tal cual.
    assert [hoja["A2"].value, hoja["F2"].value, hoja["AB2"].value, hoja["F3"].value] == \
        ["BRP01", "INV-1", "P1", "INV-2"]
    # Sigue siendo un .xlsm con macros.
    assert "xl/vbaProject.bin" in zipfile.ZipFile(ruta).namelist()


def test_el_excel_del_usuario_conserva_la_columna_ac(tmp_path, carga_sap):
    solicitud = excel_usuario(tmp_path, [FILA])
    flujo3_sap.preparar_archivo_sap(solicitud)

    hoja = load_workbook(solicitud.excel_path).worksheets[0]
    assert hoja["AC1"].value == "TXT.NUM.PRAL.AF (Nombre y NIT del acreedor)"
    assert hoja["AC2"].value == "PROVEEDOR SAS NIT 900123456"


def test_lo_que_esta_despues_de_ac_se_quita_de_la_copia(tmp_path, carga_sap):
    # El usuario dejó notas a la derecha de la plantilla (AD con encabezado,
    # AE sin encabezado): no deben viajar a SAP.
    solicitud = excel_usuario(tmp_path, [dict(FILA, AD="nota fila 2", AE="otra nota")])
    libro = load_workbook(solicitud.excel_path, keep_vba=True)
    libro.worksheets[0]["AD1"] = "NOTAS"
    libro.save(solicitud.excel_path)

    ruta = flujo3_sap.preparar_archivo_sap(solicitud)

    assert encabezados(ruta) == encabezados(PLANTILLA_SAP)
    assert load_workbook(ruta).worksheets[0].max_column == 28
    # El Excel del usuario conserva sus notas.
    hoja_usuario = load_workbook(solicitud.excel_path).worksheets[0]
    assert (hoja_usuario["AD2"].value, hoja_usuario["AE2"].value) == ("nota fila 2", "otra nota")


def test_si_falta_la_columna_a_quitar_no_queda_copia(tmp_path, carga_sap):
    # (No debería pasar: la validación ya exige AC. Es una segunda barrera.)
    ruta = tmp_path / "PDA-7889_brp_creacion.xlsm"
    shutil.copy(PLANTILLA_SAP, ruta)
    solicitud = Solicitud(case_id="PDA-7889", tipo="brp", accion="creacion", excel_path=str(ruta))

    with pytest.raises(SapError, match="no se pudieron quitar"):
        flujo3_sap.preparar_archivo_sap(solicitud)
    assert os.listdir(carga_sap) == []
