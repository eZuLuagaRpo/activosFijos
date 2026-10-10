"""
Pruebas de la validación de la plantilla "Modificación BRP" (AS02).

Se parte de la plantilla oficial (plantillas/BRP/Plantilla Modificacion
BRP.xlsm). Esa plantilla trae 1712 filas de datos: las pruebas trabajan sobre
una COPIA a la que se le borran, para escribir filas controladas.
"""

import os
import shutil

import pytest
from openpyxl import load_workbook

from core.exceptions import PlantillaInvalidaError
from core.models import Solicitud
from flujos import flujo2_validar
from validacion import router
from validacion.plantillas.brp_modificacion import ValidadorBrpModificacion

CARPETA_BRP = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "plantillas", "BRP",
)
PLANTILLA = os.path.join(CARPETA_BRP, "Plantilla Modificacion BRP.xlsm")
PLANTILLA_CREACION = os.path.join(CARPETA_BRP, "Plantilla Creación Activos BRP usuario.xlsm")

pytestmark = pytest.mark.skipif(
    not os.path.exists(PLANTILLA), reason="No está la plantilla de modificación BRP"
)


def excel_modificacion(tmp_path, filas, nombre="mod.xlsm"):
    """Copia de la plantilla oficial SIN sus filas de datos, con `filas` desde la 2."""
    destino = tmp_path / nombre
    shutil.copy(PLANTILLA, destino)
    libro = load_workbook(destino, keep_vba=True)
    hoja = libro.worksheets[0]
    hoja.delete_rows(2, hoja.max_row)
    for numero, fila in enumerate(filas, start=2):
        for letra, valor in fila.items():
            hoja[f"{letra}{numero}"] = valor
    libro.save(destino)
    return str(destino)


def validar(ruta):
    return ValidadorBrpModificacion().validar(ruta)


FILA = {"A": "7087147", "B": "0", "C": 1000, "X": "RO"}


def test_el_router_conoce_la_modificacion_brp():
    assert isinstance(router.resolver("brp", "modificacion"), ValidadorBrpModificacion)


def test_la_plantilla_oficial_es_valida_tal_cual():
    r = validar(PLANTILLA)
    assert r.valida and not r.errores_plantilla and len(r.filas) == 1712


def test_fila_con_a_y_c_es_valida(tmp_path):
    r = validar(excel_modificacion(tmp_path, [FILA, dict(FILA, A="7087188")]))
    assert r.valida and len(r.filas) == 2


def test_b_subnumero_es_opcional(tmp_path):
    assert validar(excel_modificacion(tmp_path, [dict(FILA, B=None)])).valida


@pytest.mark.parametrize("columna, texto", [
    ("A", "A (Código SAP del activo)"),
    ("C", "C (Sociedad)"),
])
def test_a_y_c_son_obligatorias(tmp_path, columna, texto):
    r = validar(excel_modificacion(tmp_path, [dict(FILA, **{columna: None})]))
    assert not r.valida
    assert texto in r.filas[0].errores[0]


def test_los_demas_campos_pueden_ir_vacios(tmp_path):
    assert validar(excel_modificacion(tmp_path, [{"A": "7087147", "C": 1000}])).valida


def test_plantilla_de_creacion_no_sirve_para_modificar():
    r = validar(PLANTILLA_CREACION)
    assert not r.valida
    assert "Faltan encabezados" in r.errores_plantilla[0]


def test_todo_o_nada_en_el_flujo_2(tmp_path):
    ruta = excel_modificacion(tmp_path, [FILA, dict(FILA, C=None)])
    sol = Solicitud(case_id="PDA-1", tipo="brp", accion="modificacion", excel_path=ruta)
    with pytest.raises(PlantillaInvalidaError) as error:
        flujo2_validar.validar_solicitud(sol)
    assert [f.fila for f in error.value.resultado.filas_invalidas] == [3]
