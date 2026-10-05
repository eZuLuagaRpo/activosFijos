"""
Pruebas de la validación de la plantilla "Creación Activos BRP".

Cada prueba parte de la PLANTILLA OFICIAL (plantillas/BRP/...xlsm), le
escribe filas de ejemplo y valida. Correr desde rpa_activos_fijos/:
    python -m pytest tests -q
"""

import os
import re
import shutil
import zipfile

import pytest
from openpyxl import load_workbook

from core.exceptions import PlantillaInvalidaError, ValidacionError
from core.models import Solicitud
from flujos import flujo2_validar
from validacion.plantillas.brp_creacion import ValidadorBrpCreacion

CARPETA_BRP = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "plantillas",
    "BRP",
)
# La plantilla que envía el USUARIO (A..AC).
PLANTILLA = os.path.join(CARPETA_BRP, "Plantilla Creación Activos BRP usuario.xlsm")

pytestmark = pytest.mark.skipif(
    not os.path.exists(PLANTILLA), reason="No está la plantilla oficial de BRP"
)


def fila_valida(**cambios):
    """Una fila con todos los obligatorios diligenciados (letra -> valor)."""
    fila = {
        "A": "BRP01",
        "B": "1000",
        "C": 1,
        "D": "COMPUTADOR PORTATIL",
        "F": "INV-0001",
        "H": "CC12345",
        "O": "NUEVO",
        "W": "AGR01",
    }
    fila.update(cambios)
    return fila


def crear_excel(tmp_path, filas, nombre="caso.xlsm", insertar_columna=False):
    """Copia la plantilla oficial y escribe `filas` desde la fila 2."""
    destino = tmp_path / nombre
    shutil.copy(PLANTILLA, destino)
    libro = load_workbook(destino, keep_vba=True)
    hoja = libro.worksheets[0]
    if insertar_columna:
        hoja.insert_cols(1)
    for numero, fila in enumerate(filas, start=2):
        for letra, valor in (fila or {}).items():
            hoja[f"{letra}{numero}"] = valor
    libro.save(destino)
    return str(destino)


def validar(ruta):
    return ValidadorBrpCreacion().validar(ruta)


# -- Casos válidos -------------------------------------------------------------

def test_plantilla_valida_varias_filas(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(), fila_valida(F="INV-0002")]))
    assert r.valida
    assert len(r.filas) == 2


@pytest.mark.parametrize("cantidad", [1, 1.0, "1", " 1 "])
def test_v2_cantidad_uno_aceptada(tmp_path, cantidad):
    r = validar(crear_excel(tmp_path, [fila_valida(C=cantidad)]))
    assert r.valida


def test_filas_vacias_intermedias_se_ignoran(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(), None, fila_valida()]))
    assert r.valida
    assert [f.fila for f in r.filas] == [2, 4]


def test_rango_usado_desactualizado_no_pierde_filas(tmp_path):
    # La plantilla oficial trae guardado el rango usado "A1:AB1". Si el archivo
    # llega así (lo guardó una herramienta que no lo actualiza), igual se deben
    # leer todas las filas y con su número real.
    ruta = crear_excel(tmp_path, [fila_valida(), None, fila_valida(C=2)])
    ruta_vieja = str(tmp_path / "rango_viejo.xlsm")
    with zipfile.ZipFile(ruta) as origen, zipfile.ZipFile(ruta_vieja, "w") as destino:
        for item in origen.infolist():
            datos = origen.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                datos = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:AB1"', datos)
            destino.writestr(item, datos)
    r = validar(ruta_vieja)
    assert [f.fila for f in r.filas] == [2, 4]
    assert [f.fila for f in r.filas_invalidas] == [4]


def test_celdas_fuera_de_a_ab_no_crean_filas(tmp_path):
    ruta = crear_excel(tmp_path, [fila_valida()])
    libro = load_workbook(ruta, keep_vba=True)
    libro.worksheets[0]["AD5"] = "nota del usuario"
    libro.save(ruta)
    r = validar(ruta)
    assert r.valida
    assert [f.fila for f in r.filas] == [2]


def test_xlsx_sin_macros_tambien_se_valida(tmp_path):
    # Desde 2026-10-04 SAP no recibe el archivo: se acepta .xlsx y .xlsm.
    ruta = crear_excel(tmp_path, [fila_valida()], nombre="caso.xlsm")
    libro = load_workbook(ruta)
    ruta_xlsx = str(tmp_path / "caso.xlsx")
    libro.save(ruta_xlsx)
    assert validar(ruta_xlsx).valida


# -- V1 · Obligatorios ---------------------------------------------------------

@pytest.mark.parametrize("vacio", [None, "", "   ", 0, "0"])
def test_v1_vacio_incluye_espacios_y_cero(tmp_path, vacio):
    r = validar(crear_excel(tmp_path, [fila_valida(D=vacio)]))
    assert not r.valida
    assert "D (Denominación)" in r.filas[0].errores[0]


def test_v1_reporta_todas_las_columnas_faltantes(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(F=None, H=None)]))
    error = r.filas[0].errores[0]
    assert "F (Número de Inventario)" in error and "H (Centro de Coste)" in error


def test_informativos_con_obligatorio_en_encabezado_pueden_ir_vacios(tmp_path):
    # E, I, P, Q, R dicen "OBLIGATORIO" en el encabezado pero son informativos.
    r = validar(crear_excel(tmp_path, [fila_valida()]))
    assert r.valida


# -- V2 · Cantidad = 1 ---------------------------------------------------------

@pytest.mark.parametrize("cantidad", [2, 5, 1.5, "uno", "2"])
def test_v2_cantidad_distinta_de_uno_invalida(tmp_path, cantidad):
    r = validar(crear_excel(tmp_path, [fila_valida(C=cantidad)]))
    assert not r.valida
    assert "debe ser 1" in r.filas[0].errores[0]


# -- V3 · Estructura -----------------------------------------------------------

def test_v3_sin_filas_de_datos(tmp_path):
    r = validar(crear_excel(tmp_path, []))
    assert not r.valida
    assert "no tiene filas de datos" in r.errores_plantilla[0]


def test_v3_encabezado_faltante(tmp_path):
    ruta = crear_excel(tmp_path, [fila_valida()])
    libro = load_workbook(ruta, keep_vba=True)
    libro.worksheets[0]["W1"] = "OTRA COSA"
    libro.save(ruta)
    r = validar(ruta)
    assert not r.valida
    assert "CLAVE AGRUPAMIENTO" in r.errores_plantilla[0]


def test_v3_columnas_corridas_se_ubican_por_encabezado(tmp_path):
    # Se inserta una columna al inicio: todo queda corrido una letra.
    ruta = crear_excel(tmp_path, [], insertar_columna=True)
    libro = load_workbook(ruta, keep_vba=True)
    hoja = libro.worksheets[0]
    for letra, valor in {"B": "BRP01", "C": "1000", "D": 1, "E": "PC", "G": "INV",
                         "I": "CC1", "P": "NUEVO", "X": "AGR01"}.items():
        hoja[f"{letra}2"] = valor
    libro.save(ruta)
    r = validar(ruta)
    assert r.valida
    assert r.advertencias_plantilla  # se reporta que las columnas están corridas


def test_formato_no_soportado(tmp_path):
    ruta = tmp_path / "caso.xls"
    ruta.write_bytes(b"")
    r = validar(str(ruta))
    assert not r.valida
    assert "debe ser .xlsx o .xlsm" in r.errores_plantilla[0]


# -- V4 · Vehículos ------------------------------------------------------------

def test_v4_vehiculo_sin_modelo_invalido(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(M="ABC123")]))
    assert not r.valida
    assert "vehículo" in r.filas[0].errores[0]


def test_v4_vehiculo_con_modelo_valido(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(M="ABC123", U="SPARK GT")]))
    assert r.valida


# -- V5 · Largos (solo advertencia) --------------------------------------------

def test_v5_largos_excedidos_son_advertencia(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(T="X" * 31, U="Y" * 16)]))
    assert r.valida
    assert len(r.filas[0].advertencias) == 2


# -- Columna extra AC (TXT.NUM.PRAL.AF) ------------------------------------------

def test_ac_vacia_es_valida(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(AC=None)]))
    assert r.valida


def test_ac_con_50_caracteres_es_valida(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(AC="X" * 50)]))
    assert r.valida and not r.filas[0].advertencias


def test_ac_con_mas_de_50_caracteres_es_error(tmp_path):
    r = validar(crear_excel(tmp_path, [fila_valida(AC="X" * 51)]))
    assert not r.valida
    assert "AC (Nombre y NIT del acreedor) tiene 51 caracteres (máximo 50)" in r.filas[0].errores[0]


def test_plantilla_sin_columna_ac_se_rechaza(tmp_path):
    # La plantilla de SAP (A..AB) no sirve como plantilla del usuario: le falta AC.
    # Se parte de la plantilla del usuario y se le borra la columna AC.
    destino = crear_excel(tmp_path, [fila_valida()], nombre="sin_ac.xlsm")
    libro = load_workbook(destino, keep_vba=True)
    libro.worksheets[0].delete_cols(29)   # AC
    libro.save(destino)
    r = validar(destino)
    assert not r.valida
    assert "AC 'TXT.NUM.PRAL.AF (Nombre y NIT del acreedor)'" in r.errores_plantilla[0]


def test_columna_s_acreedor_sigue_siendo_otra_columna(tmp_path):
    # S "ACREEDOR" e AC son campos distintos: S sin límite de 50, AC sí.
    r = validar(crear_excel(tmp_path, [fila_valida(S="Y" * 80, AC="Z" * 10)]))
    assert r.valida


# -- Regla final · Todo o nada + integración con el Flujo 2 ---------------------

def test_todo_o_nada_una_fila_mala_invalida_la_plantilla(tmp_path):
    ruta = crear_excel(tmp_path, [fila_valida(), fila_valida(C=3), fila_valida()])
    solicitud = Solicitud(case_id="PDA-1", tipo="brp", accion="creacion", excel_path=ruta)
    with pytest.raises(PlantillaInvalidaError) as error:
        flujo2_validar.validar_solicitud(solicitud)
    assert [f.fila for f in error.value.resultado.filas_invalidas] == [3]


def test_flujo2_combinacion_sin_validador(tmp_path):
    ruta = crear_excel(tmp_path, [fila_valida()])
    solicitud = Solicitud(case_id="PDA-1", tipo="prj", accion="creacion", excel_path=ruta)
    with pytest.raises(ValidacionError, match="Aún no hay validación"):
        flujo2_validar.validar_solicitud(solicitud)
