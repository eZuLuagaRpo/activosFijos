"""
Pruebas del Flujo 3 con el formulario de AS01 (BRP - Creación), usando un SAP
FALSO que anota cada acción del bot. Lo más importante que se verifica:
  - recorre fila por fila y llena los campos en el orden configurado;
  - celdas vacías no se tocan; la casilla solo se marca si hay valor;
  - NUNCA presiona Guardar si no está permitido (y nunca en MODO PRUEBAS);
  - en supervisión pide "Continuar" y sale SIN guardar;
  - si una fila falla, anota el error y CONTINÚA con la siguiente;
  - con guardar permitido: lee el código ("El act.fj. 7129560 0 se ha
    creado" -> 7129560) y escribe la columna "Código SAP".
"""

import os
import shutil
from datetime import datetime

import pytest
from openpyxl import load_workbook

import config
from core.exceptions import SapError
from core.models import Solicitud
from flujos import flujo3_sap

PLANTILLA = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "plantillas", "BRP", "Plantilla Creación Activos BRP usuario.xlsm",
)
FORMULARIO = config.SAP_FORMULARIO_AS01_BRP
GUARDAR = config.SAP_XPATH_BOTON_GUARDAR
ATRAS = config.SAP_XPATH_BOTON_ATRAS
CONFIRMAR = config.SAP_XPATH_CONFIRMAR_SALIR_SIN_GUARDAR
MENSAJE = ['//*[@id="BARRA-DE-MENSAJES"]']


def xpath_de(columna):
    return next(p["xpath"] for p in FORMULARIO if p.get("columna") == columna)


# -- SAP falso ---------------------------------------------------------------------

class Objeto:
    def __init__(self, sap, nombre, texto=""):
        self.sap, self.nombre, self.text = sap, nombre, texto

    def click(self):
        self.sap.acciones.append(("clic", self.nombre))


class FakeSap:
    """Imita SapWebGui. `mensajes`: textos que devuelve la barra tras cada
    Guardar (en orden). `fallar_en`: XPath que "no aparece" (lanza SapError)."""

    def __init__(self, mensajes=None, fallar_en=None, confirmar_visible=True):
        self.acciones = []
        self.mensajes = list(mensajes or [])
        self.fallar_en = fallar_en or []
        self.confirmar_visible = confirmar_visible

    def _revisar(self, xpaths):
        if xpaths == self.fallar_en:
            raise SapError("no apareció el campo")

    def asegurar_sesion(self):
        self.acciones.append(("sesion",))

    def recargar(self):
        self.acciones.append(("recargar",))

    def ir_a_transaccion(self, codigo):
        self.acciones.append(("transaccion", codigo))

    def enter(self):
        self.acciones.append(("enter",))

    def doble_clic_lento(self, xpaths, descripcion, espera):
        self._revisar(xpaths)
        self.acciones.append(("doble_clic", xpaths[0], espera))

    def clic(self, xpaths, descripcion):
        self._revisar(xpaths)
        self.acciones.append(("clic", xpaths[0]))

    def escribir(self, xpaths, texto, descripcion, registrar_valor=True):
        self._revisar(xpaths)
        self.acciones.append(("escribir", xpaths[0], texto))

    def existe(self, xpaths, segundos):
        if xpaths == CONFIRMAR and self.confirmar_visible:
            return Objeto(self, "confirmar_salir")
        return None

    def leer_texto(self, xpaths, descripcion):
        return self.mensajes.pop(0)


# -- Datos ------------------------------------------------------------------------

FILA = {
    "A": "BRP01", "B": 1000, "C": 1, "D": "PORTATIL", "F": "INV-1", "G": datetime(2026, 10, 5),
    "H": "CC1", "O": "WK", "W": 8000, "AC": "PROVEEDOR SAS NIT 900123456",
}


@pytest.fixture(autouse=True)
def entorno(tmp_path, monkeypatch):
    monkeypatch.setattr(flujo3_sap, "SAP_GUARDAR_REAL", False)
    monkeypatch.setattr(flujo3_sap, "MODO_PRUEBAS_REEMPLAZO", False)
    monkeypatch.setattr(flujo3_sap, "SAP_XPATH_MENSAJE_ESTADO", [])
    monkeypatch.setattr(flujo3_sap, "OUTPUT_DIR", str(tmp_path / "salidas"))


def solicitud(tmp_path, filas):
    ruta = tmp_path / "PDA-7889_brp_creacion.xlsm"
    shutil.copy(PLANTILLA, ruta)
    libro = load_workbook(ruta, keep_vba=True)
    for numero, fila in enumerate(filas, start=2):
        for letra, valor in fila.items():
            libro.worksheets[0][f"{letra}{numero}"] = valor
    libro.save(ruta)
    return Solicitud(case_id="PDA-7889", tipo="brp", accion="creacion", excel_path=str(ruta))


def permitir_guardar(monkeypatch):
    monkeypatch.setattr(flujo3_sap, "SAP_GUARDAR_REAL", True)
    monkeypatch.setattr(flujo3_sap, "SAP_XPATH_MENSAJE_ESTADO", MENSAJE)


# -- Formato de valores -------------------------------------------------------------

def test_valor_para_sap():
    assert flujo3_sap.valor_para_sap(datetime(2026, 10, 5)) == "05.10.2026"
    assert flujo3_sap.valor_para_sap(12.5) == "12,5"
    assert flujo3_sap.valor_para_sap(1.0) == "1"
    assert flujo3_sap.valor_para_sap(1000) == "1000"
    assert flujo3_sap.valor_para_sap("  CC1 ") == "CC1"


# -- Supervisión (guardar NO permitido) -----------------------------------------------

def test_supervision_llena_en_orden_pide_continuar_y_sale_sin_guardar(tmp_path):
    sap = FakeSap()
    pedidos = []
    sol = solicitud(tmp_path, [FILA, dict(FILA, F="INV-2", N="X")])

    paso = flujo3_sap.cargar_a_sap(sap, sol, esperar_continuar=pedidos.append)

    assert len(pedidos) == 2 and "fila 2" in pedidos[0] and "fila 3" in pedidos[1]
    assert ("clic", GUARDAR[0]) not in sap.acciones          # NUNCA guarda
    assert sap.acciones.count(("transaccion", "AS01")) == 2
    # Pantalla inicial en orden + Enter + Denominación + AC + Inventario + fecha formateada
    primera = sap.acciones[2:11]
    assert primera == [
        ("escribir", xpath_de("A")[0], "BRP01"),
        ("escribir", xpath_de("B")[0], "1000"),
        ("escribir", xpath_de("C")[0], "1"),
        ("enter",),
        ("escribir", xpath_de("D")[0], "PORTATIL"),
        ("escribir", xpath_de("AC")[0], "PROVEEDOR SAS NIT 900123456"),
        ("escribir", xpath_de("F")[0], "INV-1"),
        ("escribir", xpath_de("G")[0], "05.10.2026"),
        ("doble_clic", '//*[@id="M0:46:3:1::0:1-title"]', config.SAP_ESPERA_ENTRE_CLICS_SEG),
    ]
    # Celdas vacías (I, J, K...) no se tocan.
    assert not [a for a in sap.acciones if a[0] == "escribir" and a[1] == xpath_de("J")[0]]
    # La casilla N solo se marca en la fila que trae valor (una vez).
    assert sap.acciones.count(("clic", xpath_de("N")[0])) == 1
    # Sale sin guardar: Atrás (doble, lento) + confirmar.
    assert sap.acciones.count(("doble_clic", ATRAS[0], config.SAP_ESPERA_ENTRE_CLICS_SEG)) == 2
    assert sap.acciones.count(("clic", "confirmar_salir")) == 2
    # Entre filas: salir sin guardar -> RECARGAR -> /nAS01 (no antes de la 1ª).
    assert sap.acciones.count(("recargar",)) == 1
    i = sap.acciones.index(("recargar",))
    assert sap.acciones[i - 1] == ("clic", "confirmar_salir")
    assert sap.acciones[i + 1] == ("transaccion", "AS01")
    assert paso == "Revisado en SAP SIN guardar (2 fila(s))"
    assert sol.resultados_sap == {}


def test_modo_pruebas_nunca_guarda_aunque_el_interruptor_este_encendido(tmp_path, monkeypatch):
    permitir_guardar(monkeypatch)
    monkeypatch.setattr(flujo3_sap, "MODO_PRUEBAS_REEMPLAZO", True)
    sap = FakeSap()

    flujo3_sap.cargar_a_sap(sap, solicitud(tmp_path, [FILA]))

    assert ("clic", GUARDAR[0]) not in sap.acciones


def test_sin_xpath_del_mensaje_no_guarda_ni_entra_a_sap(tmp_path, monkeypatch):
    monkeypatch.setattr(flujo3_sap, "SAP_GUARDAR_REAL", True)   # pero sin SAP_XPATH_MENSAJE_ESTADO
    sap = FakeSap()
    with pytest.raises(SapError, match="No se guardó nada"):
        flujo3_sap.cargar_a_sap(sap, solicitud(tmp_path, [FILA]))
    assert sap.acciones == []


def test_error_en_una_fila_se_anota_y_continua(tmp_path):
    # La fila 2 no tiene D vacío, pero "Denominación" no aparece en SAP -> error.
    sap = FakeSap(fallar_en=xpath_de("D"))
    pedidos = []
    sol = solicitud(tmp_path, [FILA, dict(FILA, F="INV-2")])

    flujo3_sap.cargar_a_sap(sap, sol, esperar_continuar=pedidos.append)

    assert sol.resultados_sap[2].startswith("ERROR:") and sol.resultados_sap[3].startswith("ERROR:")
    assert sap.acciones.count(("transaccion", "AS01")) == 2       # siguió con la fila 3
    assert sap.acciones.count(("recargar",)) == 1                 # ...con SAP recargado
    # Tras un error NO se presiona "Atrás" (podría salir al menú / cerrar sesión).
    assert not [a for a in sap.acciones if a[:2] == ("doble_clic", ATRAS[0])]
    assert pedidos == []


# -- Guardar permitido ------------------------------------------------------------------

def test_guardar_lee_codigos_y_errores_y_escribe_la_columna(tmp_path, monkeypatch):
    permitir_guardar(monkeypatch)
    sap = FakeSap(mensajes=["El act.fj. 7129560 0 se ha creado",
                            "Centro de coste CC9 no existe"])
    sol = solicitud(tmp_path, [FILA, dict(FILA, F="INV-2", H="CC9")])

    paso = flujo3_sap.cargar_a_sap(sap, sol)

    assert sap.acciones.count(("clic", GUARDAR[0])) == 2
    assert sap.acciones.count(("recargar",)) == 1                 # también al guardar
    assert sol.resultados_sap == {2: "7129560", 3: "ERROR: Centro de coste CC9 no existe"}
    assert paso == "SAP: 1 creado(s), 1 con error"

    # Excel de respuesta: columna "Código SAP" a la derecha de AC; el original intacto.
    respuesta = tmp_path / "salidas" / "PDA-7889_brp_creacion_RESPUESTA.xlsm"
    hoja = load_workbook(respuesta).worksheets[0]
    assert (hoja["AD1"].value, hoja["AD2"].value, hoja["AD3"].value) == \
        ("Código SAP", "7129560", "ERROR: Centro de coste CC9 no existe")
    assert load_workbook(sol.excel_path).worksheets[0].max_column == 29


# -- Lectura de filas (misma lógica que la validación) -------------------------------------

def test_leer_filas_salta_vacias_y_usa_los_encabezados(tmp_path):
    from validacion.plantillas.brp_creacion import ValidadorBrpCreacion
    sol = solicitud(tmp_path, [FILA, {}, dict(FILA, F="INV-3")])
    filas = ValidadorBrpCreacion().leer_filas(sol.excel_path)
    assert [n for n, _ in filas] == [2, 4]
    assert filas[1][1]["F"] == "INV-3" and filas[0][1]["AC"] == "PROVEEDOR SAS NIT 900123456"
