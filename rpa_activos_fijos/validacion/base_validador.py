"""
validacion/base_validador.py — Motor común de validación de plantillas Excel.

Contexto de negocio: el Excel que adjunta el usuario en Appian YA ES el
formato que se carga a SAP (no se transforma). Por eso el Flujo 2 actúa como
FILTRO DE CALIDAD antes de SAP: si un registro llega incompleto, SAP lo
rechaza y el caso termina mal. Aquí se detectan esos problemas antes, fila
por fila.

Qué resuelve este motor (igual para todos los tipos de activo):
  - Abrir el Excel (.xlsx / .xlsm) y tomar la PRIMERA hoja.
  - V3 · Estructura: ubicar cada columna por su ENCABEZADO en la fila 1 (no
    por la letra) y confirmar que haya al menos una fila de datos.
  - Recorrer las filas de datos (desde la 2), ignorando las totalmente vacías.
  - V1 · Obligatorios: columnas marcadas como obligatorias con valor.
  - Largo máximo de campos informativos (solo advertencia).

Lo que es propio de cada plantilla (columnas, reglas extra como "cantidad =
1") se define en una subclase: ver validacion/plantillas/.
"""

import os
import re
from dataclasses import dataclass
from typing import Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.exceptions import ValidacionError
from core.models import ResultadoFila, ResultadoValidacion
from core.texto import normalizar

# Clasificación de cada columna dentro de la plantilla.
OBLIGATORIO = "obligatorio"      # vacío -> fila inválida
CONDICIONAL = "condicional"      # obligatorio solo si se cumple una condición (regla propia)
INFORMATIVO = "informativo"      # puede venir vacío; si viene, solo advertencias
NO_SE_VALIDA = "no_se_valida"    # no participa en las reglas


@dataclass(frozen=True)
class Columna:
    """Una columna de la plantilla tal como la entregó negocio."""

    letra: str                        # letra esperada en la plantilla oficial (ej. "C")
    encabezado: str                   # texto del encabezado en la fila 1
    campo: str                        # nombre de negocio (para los mensajes)
    tipo: str = NO_SE_VALIDA
    max_largo: Optional[int] = None   # si se excede -> advertencia...
    largo_es_error: bool = False      # ...o ERROR (invalida la fila) si es True


def normalizar_encabezado(texto):
    """
    Normaliza un encabezado para compararlo sin depender de detalles de
    tipeo: minúsculas, sin tildes, signos de puntuación y saltos de línea
    convertidos en espacio. Ej. "CTD.DE ACTIVOS FIJOS IGUALES.\\n OBLIGATORIA"
    -> "ctd de activos fijos iguales obligatoria".
    """
    texto = normalizar(texto)
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return " ".join(texto.split())


def es_vacia(valor):
    """
    Regla de negocio: una celda está VACÍA si no tiene contenido, si solo
    tiene espacios o si tiene 0 (la plantilla trata el cero como vacío).
    """
    if valor is None:
        return True
    if isinstance(valor, bool):
        return False
    if isinstance(valor, (int, float)):
        return valor == 0
    texto = str(valor).strip()
    return texto == "" or texto == "0"


def como_texto(valor):
    """Representa el valor de una celda como texto (1.0 -> "1")."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


class BaseValidador:
    """
    Clase base de los validadores de plantilla.

    Subclases deben definir:
      - nombre:   texto legible de la plantilla (ej. "BRP - Creación")
      - columnas: lista de `Columna` en el orden oficial de la plantilla
    Pueden ajustar `extensiones` (las que acepta la plantilla) y
    sobreescribir `reglas_fila()` para reglas propias.
    """

    nombre = None
    columnas = ()
    extensiones = (".xlsx", ".xlsm")

    def __init__(self, logger=None):
        self.logger = logger

    # -- API pública que usa el Flujo 2 --------------------------------------
    def validar(self, ruta_excel):
        """Valida el Excel y devuelve un `ResultadoValidacion` (nunca lanza
        por reglas de negocio; solo lanza ValidacionError si no puede abrirlo)."""
        resultado = ResultadoValidacion(archivo=ruta_excel, plantilla=self.nombre)

        _, extension = os.path.splitext(ruta_excel or "")
        if extension.lower() not in self.extensiones:
            resultado.errores_plantilla.append(
                f"Formato de archivo no soportado ('{extension or 'sin extensión'}'); "
                f"la plantilla {self.nombre} debe ser {' o '.join(self.extensiones)}."
            )
            return resultado

        filas = self._leer_primera_hoja(ruta_excel)
        encabezados = filas[0] if filas else ()

        # V3 · Estructura: ubicar columnas por encabezado.
        posiciones = self._ubicar_columnas(encabezados, resultado)
        if resultado.errores_plantilla:
            return resultado

        for numero, fila in self._filas_de_datos(filas, posiciones):
            resultado.filas.append(self._validar_fila(numero, fila))

        # V3 · Debe existir al menos una fila de datos.
        if not resultado.filas:
            resultado.errores_plantilla.append(
                "La plantilla no tiene filas de datos (desde la fila 2)."
            )
        return resultado

    def leer_filas(self, ruta_excel):
        """
        Devuelve las filas de datos como [(numero_fila, {letra_oficial: valor})],
        con la MISMA lógica de la validación (columnas por encabezado, filas
        vacías ignoradas). La usa el Flujo 3 para llenar los formularios de SAP
        con exactamente lo que se validó.

        Raises:
            ValidacionError: si no se puede abrir o faltan encabezados.
        """
        filas = self._leer_primera_hoja(ruta_excel, accion="Leyendo filas de")
        resultado = ResultadoValidacion(archivo=ruta_excel, plantilla=self.nombre)
        posiciones = self._ubicar_columnas(filas[0] if filas else (), resultado)
        if resultado.errores_plantilla:
            raise ValidacionError("; ".join(resultado.errores_plantilla))
        return list(self._filas_de_datos(filas, posiciones))

    @staticmethod
    def _filas_de_datos(filas, posiciones):
        """
        Filas de datos: desde la fila 2, ignorando las totalmente vacías. Solo
        cuentan las columnas de la plantilla: una nota fuera del rango (ej. en
        AD) no convierte la fila en un activo.
        """
        for numero, valores in enumerate(filas[1:], start=2):
            fila = {
                col.letra: (valores[indice] if indice < len(valores) else None)
                for col, indice in posiciones.items()
            }
            if all(es_vacia(v) for v in fila.values()):
                continue
            yield numero, fila

    # -- Reglas propias de cada plantilla (se sobreescribe) -------------------
    def reglas_fila(self, fila, resultado_fila):
        """
        Reglas adicionales de la plantilla. `fila` es un dict
        {letra_oficial: valor}. Agregar mensajes a resultado_fila.errores o
        resultado_fila.advertencias.
        """

    # -- Utilidades -----------------------------------------------------------
    def columna(self, letra):
        """Devuelve la `Columna` por su letra oficial."""
        return next(c for c in self.columnas if c.letra == letra)

    def etiqueta(self, letra):
        """Texto para mensajes: 'F (Número de Inventario)'."""
        return f"{letra} ({self.columna(letra).campo})"

    def _leer_primera_hoja(self, ruta_excel, accion="Validando"):
        """Devuelve todas las filas de la primera hoja como tuplas de valores."""
        try:
            # data_only: si alguna celda trae fórmula, se toma el valor calculado.
            libro = load_workbook(ruta_excel, read_only=True, data_only=True)
        except Exception as e:
            raise ValidacionError(f"No se pudo abrir el Excel '{ruta_excel}': {e}")
        try:
            hoja = libro.worksheets[0]
            # En modo read_only openpyxl confía en el rango "usado" que trae
            # guardado el archivo, y puede venir desactualizado (la plantilla
            # oficial trae A1:AB1): sin esto, se perderían las filas de datos.
            hoja.reset_dimensions()
            if self.logger:
                self.logger.info(
                    "%s '%s' (hoja '%s') con la plantilla %s.",
                    accion,
                    os.path.basename(ruta_excel),
                    hoja.title,
                    self.nombre,
                )
            return [tuple(fila) for fila in hoja.iter_rows(values_only=True)]
        finally:
            libro.close()

    def _ubicar_columnas(self, encabezados, resultado):
        """
        Busca cada columna esperada por su encabezado (no por la letra).
        - Si no aparece -> error de plantilla (V3).
        - Si aparece en otra letra -> advertencia (se valida igual, usando la
          posición real).
        Devuelve {Columna: índice_0_based}.
        """
        normalizados = [normalizar_encabezado(e) for e in encabezados]
        posiciones = {}
        faltantes = []
        for col in self.columnas:
            buscado = normalizar_encabezado(col.encabezado)
            if buscado not in normalizados:
                faltantes.append(f"{col.letra} '{col.encabezado}'")
                continue
            indice = normalizados.index(buscado)
            posiciones[col] = indice
            letra_real = get_column_letter(indice + 1)
            if letra_real != col.letra:
                resultado.advertencias_plantilla.append(
                    f"La columna '{col.encabezado}' está en {letra_real} y se "
                    f"esperaba en {col.letra}."
                )

        if faltantes:
            resultado.errores_plantilla.append(
                "Faltan encabezados en la fila 1: " + ", ".join(faltantes)
            )
        return posiciones

    def _validar_fila(self, numero, fila):
        resultado_fila = ResultadoFila(fila=numero)

        # V1 · Campos obligatorios diligenciados.
        faltantes = [
            self.etiqueta(col.letra)
            for col in self.columnas
            if col.tipo == OBLIGATORIO and es_vacia(fila.get(col.letra))
        ]
        if faltantes:
            resultado_fila.errores.append(
                "Faltan campos obligatorios: " + ", ".join(faltantes)
            )

        # Largo máximo -> advertencia (o error si la columna lo pide).
        for col in self.columnas:
            valor = fila.get(col.letra)
            if col.max_largo and not es_vacia(valor):
                largo = len(como_texto(valor))
                if largo > col.max_largo:
                    destino = (resultado_fila.errores if col.largo_es_error
                               else resultado_fila.advertencias)
                    destino.append(
                        f"{self.etiqueta(col.letra)} tiene {largo} caracteres "
                        f"(máximo {col.max_largo}); SAP podría cortarlo o rechazarlo."
                    )

        self.reglas_fila(fila, resultado_fila)
        return resultado_fila
